from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DATA_COLUMNS = ["Time", "VO2/kg", "Heart Rate", "Stage"]
LACTATE_COLUMNS = ["Stage", "Lactate"]


def parse_lactate_sheet(uploaded_file) -> list[dict]:
    raw_bytes = (
        uploaded_file.getvalue()
        if hasattr(uploaded_file, "getvalue")
        else uploaded_file.read()
    )
    file_obj = BytesIO(raw_bytes)
    xls = pd.ExcelFile(file_obj, engine="openpyxl")

    sheet_name = next(
        (
            name
            for name in xls.sheet_names
            if str(name).strip().lower() in {"lactate", "젖산", "젖산염"}
        ),
        None,
    )
    if sheet_name is None:
        return []

    file_obj.seek(0)
    df = pd.read_excel(
        file_obj,
        sheet_name=sheet_name,
        engine="openpyxl",
    )

    normalized = {
        str(col).strip().lower(): col
        for col in df.columns
    }

    stage_col = next(
        (
            normalized[key]
            for key in ["stage", "단계", "스테이지"]
            if key in normalized
        ),
        None,
    )
    lactate_col = next(
        (
            normalized[key]
            for key in [
                "lactate",
                "lactate(mmol/l)",
                "blood lactate",
                "젖산",
                "혈중 젖산염",
            ]
            if key in normalized
        ),
        None,
    )

    if stage_col is None or lactate_col is None:
        raise ValueError(
            "Lactate 시트에는 Stage와 Lactate 열이 필요합니다."
        )

    rows: list[dict] = []
    for _, row in df.iterrows():
        stage = row.get(stage_col)
        lactate = row.get(lactate_col)

        if pd.isna(stage) and pd.isna(lactate):
            continue
        if pd.isna(stage):
            continue
        if pd.isna(lactate):
            continue

        value = pd.to_numeric(lactate, errors="coerce")
        if pd.isna(value):
            raise ValueError(
                f"Lactate 시트의 '{stage}' 젖산값을 숫자로 읽을 수 없습니다."
            )

        rows.append(
            {
                "stage": str(stage).strip(),
                "lactate": float(value),
            }
        )

    return rows


def build_input_template() -> bytes:
    wb = Workbook()

    guide = wb.active
    guide.title = "작성안내"

    guide["A1"] = "NGC CPET Excel 입력 양식"
    guide["A1"].font = Font(size=16, bold=True)
    guide.merge_cells("A1:D1")

    guide_rows = [
        ("Data 시트", "운동 중 연속 측정 자료를 입력합니다."),
        ("필수 열", "Time / VO2/kg / Heart Rate / Stage"),
        (
            "Stage 작성",
            "운동 Stage에는 3, 4.5, 6 또는 100, 150, 200처럼 실제 부하값만 입력합니다. 부하 단위는 웹에서 Speed 또는 Power로 선택하며, 마지막 숫자 Stage는 자동으로 All-Out으로 처리됩니다.",
        ),
        (
            "Time 작성",
            "Excel 시간값 또는 초 단위 연속시간을 사용할 수 있습니다.",
        ),
        ("Lactate 시트", "각 Stage 종료 후 측정한 혈중 젖산염 값을 입력합니다."),
        (
            "Lactate Stage 작성",
            "Rest가 있다면 Rest를 입력하고, 각 운동강도는 Data 시트와 같은 숫자로 입력합니다. 마지막 운동강도 숫자가 All-Out이며, 이후 Recovery 측정값을 입력할 수 있습니다. Rest가 없으면 생략할 수 있습니다.",
        ),
        (
            "Lactate 작성",
            "단위는 mmol/L이며 미측정 Stage는 행을 비우거나 작성하지 않아도 됩니다.",
        ),
        (
            "주의",
            "마지막 숫자 Stage를 제외한 일반 운동 Stage는 웹에서 설정한 Stage 시간으로 자동 정규화됩니다. R 또는 Rest 구간은 웹에서 설정한 Rest 시간으로 정규화됩니다. 마지막 숫자 Stage는 All-Out으로 자동 인식해 실제 기록시간을 사용하며 Recovery도 실제 기록시간을 사용합니다.",
        ),
    ]

    row = 3
    for label, text in guide_rows:
        guide.cell(row=row, column=1, value=label)
        guide.cell(row=row, column=2, value=text)
        guide.cell(row=row, column=1).font = Font(bold=True)
        guide.cell(row=row, column=1).fill = PatternFill(
            "solid",
            fgColor="DDEBF7",
        )
        guide.merge_cells(
            start_row=row,
            start_column=2,
            end_row=row,
            end_column=4,
        )
        row += 1

    row += 1
    guide.cell(row=row, column=1, value="작성 예시")
    guide.cell(row=row, column=1).font = Font(bold=True)
    row += 1
    examples = [
        ("Data Stage 예시", "Speed: 3 → R → 4.5 → R → 6 → 7.5 / Power: 100 → R → 150 → R → 200 → 250"),
        (
            "Lactate Stage",
            "Rest → 3 → 4.5 → 6 → ... → 16 → Recovery 1min → Recovery 5min",
        ),
        (
            "Lactate",
            "0.90 → 1.10 → 1.25 → 1.55 → ... → 7.80 → 6.20 → 3.10",
        ),
    ]
    for label, text in examples:
        guide.cell(row=row, column=1, value=label)
        guide.cell(row=row, column=2, value=text)
        guide.merge_cells(
            start_row=row,
            start_column=2,
            end_row=row,
            end_column=4,
        )
        row += 1

    data = wb.create_sheet("Data")
    for col, value in enumerate(DATA_COLUMNS, start=1):
        cell = data.cell(row=1, column=col, value=value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.alignment = Alignment(horizontal="center")
    data.freeze_panes = "A2"

    lactate = wb.create_sheet("Lactate")
    for col, value in enumerate(LACTATE_COLUMNS, start=1):
        cell = lactate.cell(row=1, column=col, value=value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")
        cell.alignment = Alignment(horizontal="center")
    lactate.freeze_panes = "A2"

    for sheet in [guide, data, lactate]:
        for col_idx in range(1, sheet.max_column + 1):
            letter = get_column_letter(col_idx)
            sheet.column_dimensions[letter].width = 22
        for row_cells in sheet.iter_rows():
            for cell in row_cells:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 30
    guide.column_dimensions["C"].width = 30
    guide.column_dimensions["D"].width = 30
    data.column_dimensions["A"].width = 14
    data.column_dimensions["B"].width = 14
    data.column_dimensions["C"].width = 16
    data.column_dimensions["D"].width = 14
    lactate.column_dimensions["A"].width = 18
    lactate.column_dimensions["B"].width = 18

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
