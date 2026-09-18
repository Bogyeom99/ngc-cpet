from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def build_excel_export(report_data: dict, stage_summary, lactate_points: list[dict], lt1, lt2) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "선수정보"
    info_rows = [
        ("이름", report_data.get("name", "")),
        ("성별", report_data.get("sex", "")),
        ("구분", report_data.get("category", "")),
        ("측정날짜", report_data.get("test_date", "")),
        ("신장(cm)", report_data.get("height", "")),
        ("체중(kg)", report_data.get("weight", "")),
        ("BMI(kg/m²)", report_data.get("bmi", "")),
        ("최대 산소 섭취량(ml/kg/min)", report_data.get("vo2max", "")),
        ("최대 심박수(beats/min)", report_data.get("hrmax", "")),
        ("운동 시간", report_data.get("exercise_time", "")),
        (f"최대 {report_data.get('load_name', '')}", report_data.get("max_load", "")),
        ("단위", report_data.get("load_unit", "")),
        ("경사도(%)", report_data.get("grade_percent", "")),
        ("1페이지 평가 문구", report_data.get("page1_comment", "")),
        ("2페이지 젖산 평가 문구", report_data.get("lactate_comment", "")),
        ("LT1 분석 문구", report_data.get("lt1_comment", "")),
        ("LT2 분석 문구", report_data.get("lt2_comment", "")),
    ]
    ws.append(["항목", "값"])
    for row in info_rows:
        ws.append(list(row))

    ws2 = wb.create_sheet("스테이지요약")
    headers = [
        "Stage",
        "Load",
        "Duration_sec",
        "HR_min",
        "HR_max",
        "HR_mean",
        "VO2_min",
        "VO2_max",
        "VO2_mean",
    ]
    ws2.append(headers)
    for _, row in stage_summary.iterrows():
        ws2.append([row.get(h, "") for h in headers])

    ws3 = wb.create_sheet("젖산")
    ws3.append(["Stage", "Load", "Lactate"])
    for point in lactate_points:
        ws3.append([point.get("label", ""), point.get("load", ""), point.get("lactate", "")])

    ws4 = wb.create_sheet("LT결과")
    ws4.append(["구분", "산출여부", "부하", "사유"])
    for name, result in [("LT1", lt1), ("LT2", lt2)]:
        if result is None:
            ws4.append([name, "미선택", "", ""])
        elif result.valid:
            ws4.append([name, "산출", result.threshold_load, ""])
        else:
            ws4.append([name, "산출 불가", "", result.reason or ""])

    header_fill = PatternFill("solid", fgColor="D9E2F3")
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in sheet.columns:
            max_len = 0
            letter = col[0].column_letter
            for cell in col:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.column_dimensions[letter].width = min(max(max_len + 2, 10), 45)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
